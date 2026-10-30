"""Write the resume as an ATS-friendly Word file (single column, standard headings, no tables)."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT = "Calibri"


def _spacing(paragraph, before=0, after=2):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)


def _heading(doc, text):
    p = doc.add_paragraph()
    _spacing(p, before=8, after=3)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    # thin line under the heading
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for attr, value in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "444444")):
        bottom.set(qn(f"w:{attr}"), value)
    border.append(bottom)
    p._p.get_or_add_pPr().append(border)


def _line(doc, left, right="", bold=False):
    p = doc.add_paragraph()
    _spacing(p)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(18.5), alignment=2)  # right-aligned tab
    run = p.add_run(left)
    run.bold = bold
    if right:
        p.add_run(f"\t{right}")
    return p


def to_docx(resume, path):
    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(1.5)
        section.left_margin = section.right_margin = Cm(1.5)
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(10.5)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(name, after=0)
    run = name.add_run(resume.name)
    run.bold = True
    run.font.size = Pt(18)

    for text in (resume.headline,
                 " | ".join(v for v in resume.contact.model_dump().values() if v)):
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _spacing(p)

    _heading(doc, "Summary")
    _spacing(doc.add_paragraph(resume.summary))

    _heading(doc, "Skills")
    for group in resume.skills:
        p = doc.add_paragraph()
        _spacing(p)
        p.add_run(f"{group.category}: ").bold = True
        p.add_run(", ".join(group.items))

    _heading(doc, "Experience")
    for job in resume.experience:
        _line(doc, f"{job.title}, {job.company}", job.dates, bold=True)
        if job.location:
            doc.add_paragraph(job.location).runs[0].italic = True
        for bullet in job.bullets:
            _spacing(doc.add_paragraph(bullet, style="List Bullet"), after=1)

    if resume.projects:
        _heading(doc, "Projects")
        for project in resume.projects:
            p = doc.add_paragraph(style="List Bullet")
            _spacing(p, after=1)
            p.add_run(project.name).bold = True
            p.add_run(f": {project.description}")
            if project.link:
                p.add_run(f" ({project.link})")

    _heading(doc, "Education")
    for edu in resume.education:
        _line(doc, f"{edu.degree}, {edu.school}", edu.dates)

    doc.save(path)
    return path
